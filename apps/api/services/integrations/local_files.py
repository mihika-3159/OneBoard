import csv
import json
import io
from pathlib import Path
from typing import Dict, Any, List
import openpyxl
from pypdf import PdfReader

class LocalFilesAdapter:
    """Local file import and text extraction for CSV, XLSX, PDF, TXT, Markdown, and JSON."""
    
    @staticmethod
    def extract_text(file_bytes: bytes, filename: str) -> Dict[str, Any]:
        safe_filename = Path(filename).name
        ext = Path(filename).suffix.lower()
        
        if ext in [".txt", ".md", ""]:
            text = file_bytes.decode("utf-8", errors="replace")
            return {"file_type": ext.lstrip("."), "text": text, "metadata": {"filename": safe_filename}}
            
        elif ext == ".json":
            data = json.loads(file_bytes.decode("utf-8", errors="replace"))
            return {
                "file_type": "json",
                "text": json.dumps(data, indent=2),
                "structured_data": data,
                "metadata": {"filename": filename}
            }
            
        elif ext == ".csv":
            content = file_bytes.decode("utf-8", errors="replace")
            reader = csv.DictReader(io.StringIO(content))
            rows = list(reader)
            return {
                "file_type": "csv",
                "text": f"CSV data with {len(rows)} records. Headers: {', '.join(reader.fieldnames or [])}.\n" + content[:2000],
                "structured_data": rows,
                "metadata": {"filename": filename, "row_count": len(rows)}
            }
            
        elif ext in [".xlsx", ".xls"]:
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
            sheet = wb.active
            rows = []
            if sheet:
                for row in sheet.iter_rows(values_only=True):
                    if any(cell is not None for cell in row):
                        rows.append([str(c) if c is not None else "" for c in row])
            text_rep = "\n".join(["\t".join(r) for r in rows[:100]])
            return {
                "file_type": "xlsx",
                "text": text_rep,
                "structured_data": rows,
                "metadata": {"filename": filename, "sheet_names": wb.sheetnames, "row_count": len(rows)}
            }
            
        elif ext == ".pdf":
            reader = PdfReader(io.BytesIO(file_bytes))
            extracted_pages = []
            for idx, page in enumerate(reader.pages):
                txt = page.extract_text() or ""
                extracted_pages.append(f"--- Page {idx + 1} ---\n{txt}")
            combined = "\n\n".join(extracted_pages)
            return {
                "file_type": "pdf",
                "text": combined,
                "metadata": {"filename": filename, "page_count": len(reader.pages)}
            }
            
        else:
            raise ValueError(f"Unsupported file extension: {ext}")

local_files_adapter = LocalFilesAdapter()
